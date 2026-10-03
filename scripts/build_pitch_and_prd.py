from __future__ import annotations

import argparse
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = Path.home() / "Downloads" / "FIELDOPS AI Pitch Deck.pptx"
PRD_OUTPUT = ROOT / "docs" / "FIELDOPS_AI_PRD.docx"
DECK_OUTPUT = ROOT / "docs" / "FIELDOPS_AI_Pitch_Deck_Final.pptx"

NAVY = "183A35"
GREEN = "2F6147"
MINT = "E8F0E7"
GOLD = "D3A83D"
INK = "24332D"
MUTED = "607169"
WHITE = "FFFFFF"


def _set_cell_fill(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _set_cell_margins(cell, top=95, start=110, bottom=95, end=110) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _set_run_font(run, *, name="Aptos", size=10, color=INK, bold=False, italic=False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.bold = bold
    run.font.italic = italic
    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)


def _style_paragraph(paragraph, *, before=0, after=5, line=1.12, keep=False):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    paragraph.paragraph_format.keep_together = keep


def _add_text(document, text, *, size=10, color=INK, bold=False, italic=False, after=5, style=None):
    paragraph = document.add_paragraph(style=style)
    _style_paragraph(paragraph, after=after)
    run = paragraph.add_run(text)
    _set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return paragraph


def _add_bullet(document, text, *, level=0):
    paragraph = document.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    _style_paragraph(paragraph, after=3)
    run = paragraph.add_run(text)
    _set_run_font(run, size=9.5, color=INK)
    return paragraph


def _add_number(document, text):
    paragraph = document.add_paragraph(style="List Number")
    _style_paragraph(paragraph, after=4)
    run = paragraph.add_run(text)
    _set_run_font(run, size=9.5, color=INK)
    return paragraph


def _add_heading(document, text, level=1):
    paragraph = document.add_paragraph()
    _style_paragraph(paragraph, before=11 if level == 1 else 7, after=5, keep=True)
    run = paragraph.add_run(text)
    _set_run_font(
        run,
        size={1: 18, 2: 13, 3: 10.5}[level],
        color=NAVY if level == 1 else GREEN,
        bold=True,
    )
    return paragraph


def _add_table(document, headers, rows, widths=None, font_size=8.5):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    _set_repeat_table_header(header)
    for index, label in enumerate(headers):
        cell = header.cells[index]
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_fill(cell, NAVY)
        _set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        _style_paragraph(paragraph, after=0, line=1.0)
        run = paragraph.add_run(label)
        _set_run_font(run, size=font_size, color=WHITE, bold=True)
    for row_values in rows:
        row = table.add_row()
        for index, value in enumerate(row_values):
            cell = row.cells[index]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            _set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            _style_paragraph(paragraph, after=0, line=1.05)
            run = paragraph.add_run(str(value))
            _set_run_font(run, size=font_size, color=INK)
            if len(table.rows) % 2 == 0:
                _set_cell_fill(cell, "F3F6F1")
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def _add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("FieldOps AI  |  ")
    _set_run_font(run, size=8, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def build_prd(output_path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.3)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    header = section.header.paragraphs[0]
    header.text = "FIELDOPS AI  /  AGRICULTURE DECISION SUPPORT"
    for run in header.runs:
        _set_run_font(run, size=8, color=GREEN, bold=True)
    _add_page_number(section.footer.paragraphs[0])

    cover = document.add_paragraph()
    cover.paragraph_format.space_before = Pt(55)
    cover.paragraph_format.space_after = Pt(7)
    run = cover.add_run("FIELDOPS AI")
    _set_run_font(run, size=30, color=NAVY, bold=True)

    subtitle = document.add_paragraph()
    _style_paragraph(subtitle, after=15)
    run = subtitle.add_run("Product Requirements Document")
    _set_run_font(run, size=19, color=GREEN, bold=True)
    _add_text(document, "Urdu-first agriculture decision support for a hackathon MVP", size=12, color=MUTED, after=18)

    metadata = document.add_table(rows=0, cols=2)
    metadata.alignment = WD_TABLE_ALIGNMENT.LEFT
    metadata.autofit = False
    for label, value in [
        ("Version", "1.0 - As-built MVP"),
        ("Status", "Implemented prototype; not production agronomy software"),
        ("Primary interface", "Streamlit web app; Urdu default, English optional"),
        ("Last reviewed", date.today().isoformat()),
        ("Repository", "github.com/Anasbhatti05/FieldOps_AI_Agriculture_Multi_Agent"),
    ]:
        cells = metadata.add_row().cells
        cells[0].width = Inches(1.35)
        cells[1].width = Inches(5.9)
        for cell in cells:
            _set_cell_margins(cell, top=80, bottom=80)
        cells[0].text = ""
        cells[1].text = ""
        _set_cell_fill(cells[0], MINT)
        for cell, content, bold in ((cells[0], label, True), (cells[1], value, False)):
            paragraph = cell.paragraphs[0]
            _style_paragraph(paragraph, after=0)
            run = paragraph.add_run(content)
            _set_run_font(run, size=9, color=GREEN if bold else INK, bold=bold)

    _add_heading(document, "خلاصہ", 2)
    urdu = document.add_paragraph()
    urdu.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _style_paragraph(urdu, after=12, line=1.4)
    run = urdu.add_run(
        "فیلڈآپ اے آئی کسان کی درج کی گئی فصل، مٹی اور موسم کی معلومات کو چھ الگ جانچوں، "
        "مقامی زرعی رہنمائی اور شفاف نمونہ حساب کے ساتھ جوڑتا ہے۔ یہ مشورہ دیتا ہے، "
        "مگر پانی کی موٹر نہیں چلاتا اور تصویر سے بیماری کی حتمی تشخیص نہیں کرتا۔"
    )
    _set_run_font(run, name="Noto Nastaliq Urdu", size=11, color=INK)

    _add_heading(document, "1. Executive Summary")
    _add_text(
        document,
        "FieldOps AI is a responsive Streamlit prototype for farmers and agricultural field advisers. "
        "A user enters crop, growth stage, field area, soil moisture, recent rain, manually entered "
        "weather values, and optional notes or a crop photo. Six bounded specialist functions return "
        "structured findings. A coordinator combines them with local guidance, transparent "
        "demonstration irrigation arithmetic, and explicit human-review flags. Urdu is the default "
        "interface language; English is optional.",
    )
    _add_text(
        document,
        "This PRD describes the code as it exists today. It does not claim a live weather feed, "
        "validated crop-disease diagnosis, vector database, autonomous equipment control, or a "
        "general-purpose LLM coordinator.",
        color="8A4E25",
    )

    _add_heading(document, "2. Problem, Vision, and Goals")
    _add_text(document, "Problem: smallholder farmers and field teams must combine crop stage, soil readings, rain, weather, and visible symptoms, often without a technical interface in their preferred language.")
    _add_text(document, "Vision: make a field check understandable, evidence-aware, and reviewable before a person decides what to do.")
    _add_table(
        document,
        ["Goal", "MVP success condition"],
        [
            ("Usable for Urdu-first farmers", "Urdu is the first screen; labels, assumptions, evidence, and downloadable reports use Urdu when Urdu is selected."),
            ("Join field signals", "One submission runs six structured specialist functions over the same field record."),
            ("Make arithmetic inspectable", "The irrigation result exposes its demonstration inputs, assumptions, and calculated water volume."),
            ("Keep a human responsible", "Image uncertainty, manually entered weather, and review flags remain visible; no physical equipment is operated."),
            ("Be repeatable for a demo", "A wheat scenario, Excel test sheet, sample upload images, and automated tests are included."),
        ],
        widths=[1.55, 5.7],
    )
    _add_heading(document, "Non-goals", 2)
    for item in [
        "No live forecast provider, soil sensor integration, database, user login, or farm-history dashboard.",
        "No validated disease or pest classifier; no chemical prescription or automatic irrigation command.",
        "No claims of agronomic accuracy, market demand, water savings, or commercial readiness.",
        "No OpenAI integration. Groq vision is optional, disabled by default, and should only be enabled for a controlled private demo.",
    ]:
        _add_bullet(document, item)

    _add_heading(document, "3. Users and Jobs to Be Done")
    _add_table(
        document,
        ["User", "Need", "MVP support"],
        [
            ("Smallholder farmer", "Understand the next field check in simple Urdu.", "Urdu-first form, acre-based area, simple action list, demo workbook, downloadable farmer report."),
            ("Field adviser / agriculture technician", "Review what measurements and assumptions produced a suggestion.", "Agent findings, evidence, source filenames, calculation assumptions, review flags, JSON export."),
            ("Hackathon evaluator", "See the workflow and verify a repeatable scenario.", "Wheat demo, four Excel test cases, six-agent trace, GitHub Actions test results."),
        ],
        widths=[1.35, 2.25, 3.65],
    )

    _add_heading(document, "4. User Flow and Functional Requirements")
    for item in [
        "Choose Urdu or English; Urdu is the default.",
        "Enter crop (wheat, cotton, rice, or maize), growth stage, field area in acres, soil moisture percentage, recent rain, next-day rain estimate, temperature, wind, optional district, and optional crop notes.",
        "Optionally attach one JPG, PNG, or WEBP crop image up to 8 MB. The app resizes it for processing.",
        "Alternatively download the bilingual Excel workbook, fill a row, upload it, and select the scenario. The workbook includes a blank row, four test cases, instructions, and synthetic upload-test pictures.",
        "Run the field check; the coordinator executes six structured functions and a validator.",
        "Review the summary, action checklist, water estimate when applicable, agent evidence, calculation assumptions, confidence, and safety flags.",
        "Download a farmer-readable Markdown report or a structured JSON report. Urdu mode also localizes JSON labels and generated report text.",
    ]:
        _add_number(document, item)

    _add_heading(document, "Field Data Contract", 2)
    _add_table(
        document,
        ["Input", "Required", "Validation / source"],
        [
            ("Crop and growth stage", "Yes", "Fixed crop and stage choices; map Urdu and English workbook values."),
            ("Field area", "Yes", "Greater than zero; entered in acres, converted to hectares for calculation."),
            ("Soil moisture", "Yes", "0-100%; user reading. Demo bands are not locally calibrated."),
            ("Recent / forecast rain", "Yes", "Non-negative mm; weather is manually entered, never fetched live."),
            ("Temperature / wind", "Yes", "Temperature -10..65 C; wind 0..200 km/h; manually entered."),
            ("Location / notes", "No", "Free text. Reports exclude photo bytes; user text is not translated automatically."),
            ("Crop photo", "No", "Up to 8 MB; preliminary vision only when explicitly enabled and provider-configured."),
        ],
        widths=[1.55, 0.85, 4.85],
        font_size=8,
    )

    _add_heading(document, "5. Specialist Functions and Output Contract")
    _add_text(document, "Each function returns a structured finding with status, finding, evidence, optional calculations, confidence, limitations, recommended next action, and optional local source filename.")
    _add_table(
        document,
        ["Specialist", "Current behavior", "Boundary"],
        [
            ("Crop health / vision", "Accepts an image; without explicit private opt-in, reports that it was not sent to an AI service. Optional Groq image screening describes visible signs.", "Always preliminary; low confidence; cannot diagnose."),
            ("Pest screening", "Reports whether an image was supplied and requests a human inspection.", "No validated pest model; never recommends spraying."),
            ("Soil", "Compares entered moisture with demo screening bands (<20%, 20-26%, >=26%).", "Bands are demo defaults, not soil-calibrated."),
            ("Weather", "Summarizes entered rainfall, temperature, and wind; adds simple risk notes.", "No live weather source; low confidence."),
            ("Irrigation", "Runs deterministic, rain-adjusted demonstration arithmetic.", "Estimate only; no pump control or field-calibrated prescription."),
            ("Knowledge / retrieval", "Ranks local Markdown paragraphs using token overlap and returns matching passages/source filenames.", "Small bilingual demo corpus; not vector search or district-verified guidance."),
            ("Coordinator + validator", "Combines the six findings, prioritizes a checklist, and flags uncertainty/conflicts.", "Deterministic orchestration; not a general LLM reasoning agent."),
        ],
        widths=[1.45, 3.35, 2.45],
        font_size=7.8,
    )

    _add_heading(document, "6. Irrigation Estimate and Assumptions")
    _add_text(document, "Current demonstration formula; the numbers below are configurable code defaults, not agronomic recommendations:", bold=True)
    _add_table(
        document,
        ["Step", "Implemented calculation"],
        [
            ("Moisture gap", "max(0, 26% demo target - entered soil moisture %)"),
            ("Gross deficit", "moisture gap (percentage points) x 0.3 m demo root depth x 10 = mm"),
            ("Effective rain", "entered forecast rain (mm) x 0.80 assumed effective fraction"),
            ("Net deficit", "max(0, gross deficit - effective rain)"),
            ("Volume", "net deficit (mm) x area (hectares) x 10 = m3; acres convert at 0.404686 ha/acre"),
        ],
        widths=[1.45, 5.8],
    )
    _add_text(document, "Example test case: 5 acres of wheat (2.02343 ha), 17% entered moisture, and 0 mm forecast rain produces approximately 546.3 m3 in this demo formula (the UI rounds to 546 m3). With 40 mm forecast rain the net estimate is zero and the app asks for a field recheck. These outputs are for software demonstration, not instructions to irrigate.", color="8A4E25")

    _add_heading(document, "7. Safety, Privacy, and Deployment")
    _add_table(
        document,
        ["Risk / constraint", "MVP control"],
        [
            ("Exposed credentials", "Previously exposed keys must be revoked. Store replacements only in local ignored .env or provider secrets; never commit them."),
            ("Paid provider abuse", "GROQ_VISION_ENABLED=false by default. A key alone cannot call Groq; public app must remain disabled without auth/rate limits."),
            ("Image privacy", "Without explicit vision opt-in, image is not sent to an AI provider. With opt-in it is sent to Groq; disclose this and charges."),
            ("Wrong farm action", "Show demo assumptions, low confidence, conflicts, human verification, and no-actuator boundary."),
            ("Public personal data", "No login, access control, database, or defined retention workflow. Do not upload confidential farm records."),
            ("Hosting", "Primary path is GitHub -> Streamlit Community Cloud, entry file app.py. Render/Docker are optional alternatives."),
        ],
        widths=[1.65, 5.6],
        font_size=8,
    )

    _add_heading(document, "8. Acceptance Criteria and Evaluation")
    _add_table(
        document,
        ["Check", "Expected result"],
        [
            ("Urdu UI and report", "Default screen, result, assumptions, evidence labels, Markdown, and JSON use Urdu; regression test rejects English alphabet leakage from generated report text."),
            ("Excel workflow", "Four sample rows parse; filled blank row accepts Urdu choices/digits; invalid rows show field-level errors."),
            ("Dry wheat scenario", "5-acre/17%-moisture/no-forecast-rain test returns a repeatable demo estimate and explicit assumptions."),
            ("Rain offset scenario", "40 mm forecast offsets the demo deficit; flags a soil recheck instead of triggering irrigation."),
            ("Photo without vision", "Upload completes; no AI-provider call occurs when vision is disabled, even if a key exists."),
            ("Hosted secrets", "Mocked Streamlit secrets work when configured; missing secrets leave vision disabled."),
            ("CI and runtime", "GitHub Actions tests Python 3.11 and 3.13; local environment has 15 passing tests, clean dependency check, and healthy Streamlit endpoint."),
        ],
        widths=[1.55, 5.7],
        font_size=7.8,
    )
    _add_text(document, "Not yet measured: agronomic accuracy, disease sensitivity, retrieval relevance on expert-labeled data, field-water savings, latency at production scale, and farmer usability in supervised field studies.", italic=True, color=MUTED)

    _add_heading(document, "9. Roadmap")
    _add_table(
        document,
        ["Phase", "Candidate work", "Gate before launch"],
        [
            ("Near term", "Field-adviser review of Urdu wording, local crop guidance, and demo thresholds.", "Signed agronomist review and documented corrections."),
            ("Integration", "Trusted weather source, district/crop-specific knowledge, source/version tracking.", "Provider coverage, freshness, citations, and evaluation set."),
            ("Vision", "Crop-specific validated model and labeled Pakistan field-photo evaluation set.", "Measured false-positive/false-negative rates; human escalation."),
            ("Public product", "Login, rate limiting, consent, audit storage, deletion, and support process.", "Privacy/security review before collecting real farm data."),
        ],
        widths=[1.0, 3.6, 2.65],
        font_size=8,
    )

    _add_heading(document, "10. Release Decision")
    _add_text(document, "Suitable for: hackathon demonstration, controlled usability feedback, and software workflow testing. Not suitable for: unsupervised farm decisions, disease diagnosis, pesticide selection, paid public vision calls, or physical irrigation control.", bold=True)
    _add_text(document, "Primary live-demo path: streamlit.app public deployment with all secrets empty and Groq vision disabled.", color=GREEN)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = "FieldOps AI Product Requirements Document"
    document.core_properties.subject = "As-built agriculture decision-support hackathon MVP"
    document.core_properties.author = "FieldOps AI Project Team"
    document.save(output_path)


DECK_REPLACEMENTS = {
    1: {
        "AI-Powered Agriculture Decision Assistant": "Urdu-First Field Decision Support",
        "From Raw Field Data to Actionable, Explainable Decisions.": "One field check. Clear evidence. A human-reviewed next step.",
        "Multi-Agent RAG System": "Six Specialist Checks",
        "Deterministic Logic": "Transparent Demo Estimate",
    },
    2: {
        "Agriculture Decisions Are Data-Heavy": "Field Decisions Need Joined-Up Evidence",
        "Flowering / Yield Risk": "Growth-Stage Context",
        "Soil Moisture": "Entered Soil Moisture",
        "Sensor Thresholds": "No Live Sensor Feed",
        "Weather Forecast": "Entered Weather Values",
        "Rain & Temp Spikes": "Rain and Temperature",
        "Leaf Symptoms": "Optional Crop Photo",
        "Visual Pathology": "Preliminary Screening",
        "Collecting field data is solved. The challenge is converting fragmented inputs into immediate, reliable farm action.": "Field details arrive as separate readings. The challenge is turning them into a clear, cautious checklist without claiming certainty.",
    },
    3: {
        "Transforming Agricultural Inputs Into an Explainable Action Plan": "Combining field inputs into a reviewable action checklist",
        "Clear operational steps built on verified context and logic.": "Clear next checks built on entered data and stated assumptions.",
        "Dedicated specialists for crop health, soil, and weather analysis.": "Six bounded checks: crop photo, pest, soil, weather, irrigation, and knowledge.",
        "Grounded agronomic guidelines for evidence-based advice.": "A small local bilingual guidance set; not district-certified.",
        "Exact calculations paired with output verification safety.": "Deterministic demo math paired with visible uncertainty and human review.",
        "Specialized AI Agents": "Specialist Checks",
        "RAG Knowledge Base": "Local Guidance Search",
        "Deterministic Engine & Validation": "Demo Calculation + Review Flags",
    },
    4: {
        "Raw Field Data": "Field Entry",
        "Crop Stage": "Crop Stage",
        "Soil Sensors": "Entered Soil Moisture",
        "Weather Forecast": "Entered Weather",
        "Crop Leaf Image": "Optional Crop Photo",
        "AI Agents": "Specialist Checks",
        "Crop Health": "Crop-Photo Screen",
        "Soil Analysis": "Soil Check",
        "Weather Risk": "Weather Summary",
        "Irrigation Agent": "Pest + Water Checks",
        "RAG Context": "Local Guide Search",
        "Agronomic Rules": "Demo Guidance Notes",
        "Evidence Lookup": "Keyword-Matched Passages",
        "Field Guidelines": "Source Files Shown",
        "Engine Check": "Deterministic Estimate",
        "Water Thresholds": "Uncalibrated Defaults",
        "Validator Audit": "Review Flags",
        "Consistency Test": "Human Check Required",
        "Action Plan": "Next-Step Checklist",
        "Action Steps": "Prioritized Checks",
        "Evidence Base": "Evidence + Assumptions",
        "Confidence Rating": "Uncertainty Visible",
    },
    5: {
        "One Problem. Multiple Specialized AI Agents.": "One Coordinator. Six Specialist Functions.",
        "COORDINATOR AGENT": "FIELD COORDINATOR",
        "Orchestrates specialist execution and aggregates final reports": "Combines structured findings, adds review flags, and exports a checklist.",
        "Crop-Health Agent": "Crop-Photo Screen",
        "Analyzes leaf condition and visual symptoms.": "Optional preliminary image screening; never a confirmed diagnosis.",
        "Soil Agent": "Soil Check",
        "Evaluates soil moisture levels and sensor data.": "Checks one entered reading against demo screening bands.",
        "Weather Agent": "Weather Summary",
        "Tracks rainfall predictions and heat indexes.": "Summarizes manually entered rain, temperature, and wind.",
        "Irrigation Agent": "Irrigation Estimate",
        "Executes reasoning with agricultural rule engine.": "Runs transparent demo arithmetic; it never operates equipment.",
        "Validator Agent": "Knowledge + Validator",
        "Audits output consistency before final response.": "Searches local notes and flags uncertainty for human review.",
    },
    6: {
        "AI + Deterministic Logic": "Optional Vision + Demo Rules",
        "AI Reasoning + Rules/Calculations + Validation": "Bounded Checks + Deterministic Math + Human Review",
        "We do not rely on an LLM alone for critical numeric decisions.": "No language model calculates or authorizes the water estimate.",
        "AI Reasoning": "Optional Groq Vision",
        "Contextual understanding.": "Private opt-in only; disabled for public demos.",
        "Image symptom interpretation.": "Preliminary visible-sign screening; no diagnosis.",
        "Natural language recommendations.": "Localized checklist text; not a general chat assistant.",
        "Deterministic Rules": "Transparent Demo Math",
        "Numeric thresholds.": "26% target is an uncalibrated demo default.",
        "Reproducible irrigation calculations.": "Repeatable, rain-adjusted arithmetic.",
        "Boundary condition verification.": "Input validation plus human-review flags.",
    },
    7: {
        "RELIABILITY & GROUNDING": "LOCAL KNOWLEDGE",
        "Recommendations Backed by Knowledge": "Small Guidance Set, Visible Limits",
        "RAG EXECUTION FLOW": "LOCAL SEARCH FLOW",
        "Retrieval-Augmented Generation (RAG)": "Keyword Matching Over Local Markdown",
        "Retrieves verified agricultural manuals and guidelines before response generation.": "Searches two small demo guides; they are not district-certified manuals.",
        "1. User Input Received": "1. Field Input",
        "2. Knowledge Search": "2. Keyword Search",
        "3. Contextual AI Reasoning": "3. Matching Passages",
        "4. Evidence-Based Output": "4. Coordinator Review",
        "Supporting Evidence": "Source Files",
        "Cited references for trust.": "Filename shown; matched passages available to inspect.",
        "Transparent Output": "Known Limits",
        "Auditable decision trail.": "No vector index, live feed, or expert-validated corpus.",
    },
    8: {
        "Turning Field Conditions Into an Irrigation Decision": "Dry Wheat Walkthrough: Demo, Not Prescription",
        "Field Inputs": "Controlled Test Inputs",
        "Execute Immediate Irrigation. Moisture below critical 250 threshold during flowering yield window.": "Demo estimate: about 546 m3. Check soil near roots before any decision.",
        "Verified by Validator Engine": "Human Review Required",
        "Action Recommendation": "Suggested Next Check",
        "High Confidence": "Low Confidence",
        "2 Acres": "5 Acres",
        "220 (Low)": "17% (Entered)",
        "3 mm / 2 mm": "1 mm / 0 mm",
        "30°C": "34°C",
    },
    9: {
        "VISUAL ASSISTANCE": "OPTIONAL PHOTO REVIEW",
        "Crop Health Analysis": "Private Vision Opt-In",
        "Assistive feature to aid decision making, not an automated diagnosis.": "Uploaded photos are not sent to an AI service unless vision is explicitly enabled.",
        "IMAGE ANALYSIS AGENT": "GROQ VISION - OPTIONAL",
        "Assistive Visual Screening": "Off for Public Demos",
        "Upload leaf photo to evaluate visual stress markers and potential issues.": "Private opt-in sends the photo to Groq and may incur API charges.",
        "Observation": "Limit",
        "Early rust lesion patterns.": "No validated crop-specific diagnostic model is included.",
        "Action Step": "Required",
        "Monitor field humidity levels.": "A local adviser must verify symptoms before treatment.",
    },
    10: {
        "Simple Input. Clear Action.": "Simple Entry. Cautious Checklist.",
        "Field: North Plot #2": "Demo field: Multan",
        "Moisture: 220": "Moisture: 17%",
        "Weather: 30°C Clear": "Entered: 34°C, 0 mm forecast rain",
        "Irrigate Immediately": "Possible Moisture Gap",
        "Moisture below critical flowering threshold.": "Uncalibrated demo estimate; verify soil with a person.",
        "Evidence Included": "Assumptions Shown",
        "Trace Log Ready": "Reports Downloadable",
    },
    11: {
        "See How the Decision Was Made": "Inspect Inputs, Findings, and Assumptions",
        "Full Auditability": "Reviewable Specialist Trace",
        "Explainability is built into the product experience rather than a hidden black box.": "The app shows structured findings and review flags; it is not a production audit system.",
        "Inspectable step-by-step reasoning logs.": "Expandable evidence, calculations, and source filenames.",
        "1. Coordinator Agent Initialized": "1. Six Specialist Checks Run",
        "2. Soil & Weather Agents Processed Data": "2. Soil + Manual Weather Reviewed",
        "3. Rule Engine Irrigation Check": "3. Demo Formula Calculates Estimate",
        "4. Validator Consistency Passed": "4. Uncertainty Flags Shown for Human Review",
        "5. Final Action Plan Rendered": "5. Checklist + Reports Displayed",
    },
    12: {
        "Built With Modern AI & Web Technologies": "Built With a Small, Testable MVP Stack",
        "Python": "Python",
        "Agent Application": "Coordinator + Specialist Functions",
        "Streamlit": "Streamlit",
        "Interactive Interface": "Urdu-First Responsive UI",
        "Google Gemini": "Groq Vision (Optional)",
        "LLM Reasoning": "Private Explicit Opt-In Only",
        "RAG Framework": "Local Knowledge Search",
        "Knowledge Retrieval": "Markdown + Token Overlap",
        "Rule Engine": "Calculator Tool",
        "Deterministic Calculations": "Documented Demo Formula",
        "Validation Layer": "Validator Flags",
        "Consistency Checks": "Human Review Required",
        "GitHub / Colab": "GitHub Actions",
        "Development Workflow": "Python 3.11 + 3.13 CI",
    },
    13: {
        "Beyond a Generic Agriculture Chatbot": "A Bounded Field-Check Workflow",
        "Single model prompt": "Six focused specialist functions",
        "Mostly text conversations": "Structured field inputs + optional photo",
        "LLM-based numbers": "Deterministic demo arithmetic",
        "Limited transparency": "Visible assumptions + review flags",
        "Specialized Multi-Agents": "Six Bounded Checks",
        "Grounded RAG Evidence": "Local Keyword-Matched Passages",
        "Validation & Trace Audit": "Human Review; No Production Audit Store",
    },
    14: {
        "Live Demo Journey": "Five-Minute Demo Journey",
        "STEP 1 - 3": "STEP 1",
        "Configure Field Input": "Load Wheat Demo",
        "Select Wheat, Flowering stage, field size.": "Review five-acre field inputs.",
        "STEP 4 - 6": "STEP 2",
        "Enter Environmental Data": "Run the Field Check",
        "Add soil moisture, weather, leaf photo.": "Use entered moisture and manual weather; photo is optional.",
        "STEP 7 - 8": "STEP 3",
        "Run FIELDOPS AI": "Inspect Specialist Findings",
        "Agents execute analysis workflow.": "Six bounded checks return structured results.",
        "STEP 9": "STEP 4",
        "View Calculations": "Review Demo Calculation",
        "Review deterministic moisture checks.": "Open assumptions; explain why the estimate is not advice.",
        "STEP 10": "STEP 5",
        "Audit Evidence": "Review and Export",
        "Check RAG references and confidence.": "Show source filenames, uncertainty, and human-review flags.",
        "STEP 11": "OPTIONAL",
        "Download Action Plan": "Try Excel + Report",
        "Export operational field instructions.": "Upload a test row and download the bilingual report.",
    },
    15: {
        "Where FIELDOPS AI Can Help": "Potential Users and MVP Use Cases",
        "Primary Use Cases": "MVP-Supported Checks",
        "Irrigation decision support.": "Rain-adjusted demo water estimate.",
        "Crop-health screening.": "Preliminary image workflow; no diagnosis.",
        "Weather-aware farm planning.": "Summary of manually entered weather.",
        "Knowledge assistance.": "Search of a small local guidance set.",
    },
    16: {
        "From Hackathon MVP to Scalable Platform": "Next Steps After Expert Review",
        "Live API Data": "Trusted Weather Feed",
        "Satellite weather integrations.": "Timestamped local forecasts with source attribution.",
        "Expanded Crop Knowledge": "Local Agronomy Review",
        "Broader disease libraries.": "Validate rules and Urdu guidance by crop and district.",
        "Regional Intelligence": "Validated Crop Vision",
        "Localized microclimate tuning.": "Test on an expert-labeled local photo set.",
        "Mobile Experience": "Public-Product Readiness",
        "Offline-first farmer App.": "Add login, consent, rate limits, and low-bandwidth support.",
    },
    17: {
        "Understand the Field. Reason Across Data. Recommend the Next Action.": "Read the field evidence. / Show assumptions and uncertainty. / Leave decisions to people.",
        "AI Agents + RAG + Deterministic Logic + Validation": "Six Checks + Local Search + Demo Math + Human Review",
    },
    18: {
        "HACKATHON PRESENTATION": "FIELDOPS AI  /  HACKATHON MVP",
        "Questions & Feedback": "Questions & Feedback",
    },
    19: {
        "Image Sources": "Image Sources and Attributions",
    },
}


def _normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def _replace_shape_text(shape, replacement: str) -> None:
    text_frame = shape.text_frame
    paragraphs = text_frame.paragraphs
    first_paragraph = paragraphs[0]
    first_run = first_paragraph.runs[0] if first_paragraph.runs else first_paragraph.add_run()
    first_run.text = replacement
    for run in first_paragraph.runs[1:]:
        run.text = ""
    for paragraph in paragraphs[1:]:
        for run in paragraph.runs:
            run.text = ""
    text_frame.word_wrap = True
    text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


def revise_pitch_deck(source: Path, output_path: Path) -> None:
    presentation = Presentation(source)
    replaced = set()
    for slide_no, slide in enumerate(presentation.slides, 1):
        replacements = DECK_REPLACEMENTS.get(slide_no, {})
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False) or not shape.text.strip():
                continue
            old_text = _normalize_text(shape.text)
            if old_text in replacements:
                _replace_shape_text(shape, replacements[old_text])
                replaced.add((slide_no, old_text))

    missing = [
        f"slide {slide_no}: {old_text}"
        for slide_no, replacements in DECK_REPLACEMENTS.items()
        for old_text in replacements
        if (slide_no, old_text) not in replaced
    ]
    if missing:
        raise ValueError("Expected slide text not found:\n" + "\n".join(missing))

    all_text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )
    for unsupported_claim in [
        "Execute Immediate Irrigation",
        "250 threshold",
        "Google Gemini",
        "Early rust lesion patterns",
        "Verified by Validator Engine",
    ]:
        if unsupported_claim in all_text:
            raise ValueError(f"Unsupported claim remains in deck: {unsupported_claim}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the as-built FieldOps AI PRD and updated pitch deck.")
    parser.add_argument("--source-pptx", type=Path, default=DEFAULT_SOURCE, help="Original user-provided pitch deck")
    parser.add_argument("--prd-output", type=Path, default=PRD_OUTPUT)
    parser.add_argument("--deck-output", type=Path, default=DECK_OUTPUT)
    args = parser.parse_args()
    if not args.source_pptx.exists():
        raise SystemExit(f"Source deck not found: {args.source_pptx}. Pass --source-pptx with its location.")
    build_prd(args.prd_output)
    revise_pitch_deck(args.source_pptx, args.deck_output)
    print(f"PRD: {args.prd_output}")
    print(f"Deck: {args.deck_output}")


if __name__ == "__main__":
    main()
