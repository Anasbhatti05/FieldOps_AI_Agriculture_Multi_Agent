import unittest
import io
from unittest.mock import patch

from openpyxl import load_workbook

from agents.coordinator import coordinate
from tools.spreadsheet import create_test_workbook, read_workbook


class SpreadsheetWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.scenarios = read_workbook(create_test_workbook())

    def test_template_has_blank_entry_and_four_examples(self):
        self.assertEqual(len(self.scenarios), 5)
        self.assertIsNone(self.scenarios[0].field)
        self.assertIn("فصل", self.scenarios[0].error)
        self.assertEqual(self.scenarios[1].field.crop, "Wheat")
        self.assertEqual(self.scenarios[1].field.soil_moisture_pct, 17)

    def test_workbook_embeds_three_upload_test_pictures(self):
        workbook = load_workbook(io.BytesIO(create_test_workbook()))

        self.assertIn("Sample Photos", workbook.sheetnames)
        self.assertEqual(len(workbook["Sample Photos"]._images), 3)

    def test_spreadsheet_cases_produce_different_irrigation_outcomes(self):
        dry_field = self.scenarios[1].field
        rain_field = self.scenarios[2].field
        with patch.dict("os.environ", {"GROQ_API_KEY": ""}):
            dry_result = coordinate(dry_field)
            rain_result = coordinate(rain_field)

        dry_water = next(item for item in dry_result.findings if item.agent == "Irrigation agent").calculation["estimated_net_water_m3"]
        rain_water = next(item for item in rain_result.findings if item.agent == "Irrigation agent").calculation["estimated_net_water_m3"]
        self.assertGreater(dry_water, 0)
        self.assertEqual(rain_water, 0)
        self.assertTrue(any("forecast rain" in flag for flag in rain_result.review_flags))

    def test_filled_blank_row_accepts_urdu_choices_and_digits(self):
        workbook = load_workbook(io.BytesIO(create_test_workbook()))
        sheet = workbook["Field Data"]
        for cell, value in {
            "B2": "گندم",
            "C2": "پھول یا بالی",
            "D2": "۵",
            "E2": "۱۷",
            "F2": "۱",
            "G2": "۰",
            "H2": "۳۴",
            "I2": "۸",
        }.items():
            sheet[cell] = value
        output = io.BytesIO()
        workbook.save(output)
        scenario = read_workbook(output.getvalue())[0]

        self.assertEqual(scenario.field.crop, "Wheat")
        self.assertEqual(scenario.field.field_area_acres, 5)
        self.assertEqual(scenario.field.soil_moisture_pct, 17)


if __name__ == "__main__":
    unittest.main()