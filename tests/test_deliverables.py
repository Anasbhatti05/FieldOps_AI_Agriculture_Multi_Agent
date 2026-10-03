from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.util import Inches

from scripts.build_pitch_and_prd import DECK_REPLACEMENTS, revise_pitch_deck


ROOT = Path(__file__).resolve().parents[1]
PRD_PATH = ROOT / "docs" / "FIELDOPS_AI_PRD.docx"
DECK_PATH = ROOT / "docs" / "FIELDOPS_AI_Pitch_Deck_Final.pptx"


def _deck_text(path: Path) -> str:
    presentation = Presentation(path)
    return "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )


def _prd_text() -> str:
    document = Document(PRD_PATH)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    paragraphs.extend(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    return "\n".join(paragraphs)


def test_prd_covers_as_built_scope_and_demo_formula():
    content = _prd_text()

    assert "Product Requirements Document" in content
    assert "26% demo target" in content
    assert "token overlap" in content
    assert "does not claim a live weather feed" in content
    assert "not suitable for" in content.lower()
    assert "15 passing tests" in content
    assert "approximately 546.3 m3" in content


def test_pitch_rewriter_corrects_claims_in_a_synthetic_deck(tmp_path):
    source = Presentation()
    for slide_number in range(1, 20):
        slide = source.slides.add_slide(source.slide_layouts[6])
        for original_text in DECK_REPLACEMENTS.get(slide_number, {}):
            textbox = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(0.4))
            textbox.text = original_text
    source_path = tmp_path / "synthetic_pitch_source.pptx"
    output_path = tmp_path / "synthetic_pitch_revised.pptx"
    source.save(source_path)

    revise_pitch_deck(source_path, output_path)
    presentation = Presentation(output_path)
    content = _deck_text(output_path)

    assert len(presentation.slides) == 19
    assert "about 546 m3" in content
    assert "250 threshold" not in content
    assert "Execute Immediate Irrigation" not in content
    assert "Google Gemini" not in content
    assert "Early rust lesion patterns" not in content
    assert "never a confirmed diagnosis" in content