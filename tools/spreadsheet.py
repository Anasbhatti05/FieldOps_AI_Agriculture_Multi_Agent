from __future__ import annotations

import io
import math
import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as SpreadsheetImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from agents.models import FieldAssessment


HEADERS = [
    "Test case / جانچ کا نام",
    "Crop / فصل",
    "Growth stage / فصل کی حالت",
    "Field area (acres) / رقبہ (ایکڑ)",
    "Soil moisture (%) / مٹی کی نمی (فیصد)",
    "Recent rain (mm) / پچھلی بارش (ملی میٹر)",
    "Forecast rain (mm) / متوقع بارش (ملی میٹر)",
    "Temperature (C) / درجہ حرارت (سینٹی گریڈ)",
    "Wind (km/h) / ہوا کی رفتار (کلومیٹر فی گھنٹہ)",
    "Location / ضلع",
    "Crop notes / فصل میں تبدیلی",
]

HEADER_ALIASES = {
    "scenario": ("test case / جانچ کا نام", "test case", "scenario", "جانچ کا نام"),
    "crop": ("crop / فصل", "crop", "فصل"),
    "stage": ("growth stage / فصل کی حالت", "growth stage", "crop stage", "فصل کی حالت", "فصل کا مرحلہ"),
    "area": ("field area (acres) / رقبہ (ایکڑ)", "field area acres", "area acres", "رقبہ ایکڑ", "رقبہ"),
    "moisture": ("soil moisture (%) / مٹی کی نمی (فیصد)", "soil moisture", "soil moisture pct", "مٹی کی نمی", "نمی فیصد"),
    "recent_rain": ("recent rain (mm) / پچھلی بارش (ملی میٹر)", "recent rain mm", "recent rainfall", "پچھلی بارش", "گزشتہ بارش"),
    "forecast_rain": ("forecast rain (mm) / متوقع بارش (ملی میٹر)", "forecast rain mm", "expected rain", "متوقع بارش", "اگلے دن کی بارش"),
    "temperature": ("temperature (c) / درجہ حرارت (سینٹی گریڈ)", "temperature c", "temperature", "درجہ حرارت"),
    "wind": ("wind (km/h) / ہوا کی رفتار (کلومیٹر فی گھنٹہ)", "wind km h", "wind speed", "ہوا کی رفتار"),
    "location": ("location / ضلع", "location", "district", "ضلع", "علاقہ"),
    "notes": ("crop notes / فصل میں تبدیلی", "crop notes", "notes", "فصل میں تبدیلی", "نوٹ"),
}

FIELD_LABELS_UR = {
    "crop": "فصل",
    "stage": "فصل کی حالت",
    "area": "رقبہ",
    "moisture": "مٹی کی نمی",
    "recent_rain": "پچھلی بارش",
    "forecast_rain": "متوقع بارش",
    "temperature": "درجہ حرارت",
    "wind": "ہوا کی رفتار",
}

REQUIRED_FIELDS = tuple(FIELD_LABELS_UR)
URDU_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


@dataclass(frozen=True)
class SpreadsheetScenario:
    name: str
    sheet_name: str
    row_number: int
    field: FieldAssessment | None
    error: str = ""


def _normalize_header(value: object) -> str:
    return re.sub(r"[^\w]+", " ", str(value or "").lower()).strip()


NORMALIZED_ALIASES = {
    _normalize_header(alias): key
    for key, aliases in HEADER_ALIASES.items()
    for alias in aliases
}


def create_test_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Field Data"
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "B2"
    sheet.append(HEADERS)
    sheet.append(["میرا کھیت / My field"] + [None] * (len(HEADERS) - 1))
    sheet.append([
        "خشک گندم / Dry wheat",
        "گندم",
        "پھول یا بالی",
        5,
        17,
        1,
        0,
        34,
        8,
        "Multan",
        "کچھ پتے پیلے دکھائی دے رہے ہیں",
    ])
    sheet.append([
        "بارش کے بعد جانچ / Rain offset",
        "گندم",
        "پھول یا بالی",
        5,
        17,
        1,
        40,
        34,
        8,
        "Multan",
        "بارش کے بعد مٹی دوبارہ چیک کریں",
    ])
    sheet.append([
        "گرمی اور تیز ہوا / Heat and wind",
        "کپاس",
        "بڑھوتری",
        3,
        18,
        0,
        0,
        40,
        30,
        "Bahawalpur",
        "گرمی کے دباؤ اور تیز ہوا کی جانچ کریں",
    ])
    sheet.append([
        "زیادہ نمی / Moisture check",
        "چاول",
        "بڑھوتری",
        2,
        30,
        8,
        12,
        29,
        10,
        "Larkana",
        "عام نگرانی کا نمونہ",
    ])
    sheet.auto_filter.ref = f"A1:{chr(64 + len(HEADERS))}{sheet.max_row}"

    header_fill = PatternFill("solid", fgColor="28583E")
    header_font = Font(name="Noto Sans Arabic", color="FFFFFF", bold=True, size=11)
    thin_green = Side(style="thin", color="D5E0D2")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="D3A83D"))
    sheet.row_dimensions[1].height = 52
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        for cell in row:
            cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin_green)
            if cell.row == 2:
                cell.fill = PatternFill("solid", fgColor="FFF5D7")
    for index, width in enumerate([28, 22, 28, 23, 25, 27, 28, 30, 34, 22, 38], start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row_number in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_number].height = 38

    crop_validation = DataValidation(type="list", formula1='"گندم,کپاس,چاول,مکئی"', allow_blank=True)
    stage_validation = DataValidation(
        type="list",
        formula1='"ابتدائی پودا,بڑھوتری,پھول یا بالی,فصل پکنے کے قریب"',
        allow_blank=True,
    )
    sheet.add_data_validation(crop_validation)
    sheet.add_data_validation(stage_validation)
    crop_validation.add("B2:B100")
    stage_validation.add("C2:C100")

    guide = workbook.create_sheet("Instructions")
    guide.sheet_view.rightToLeft = True
    guide.append(["Excel Field Test / ایکسل میں کھیت کی جانچ"])
    guide.append(["1. 'Field Data' sheet میں پیلی 'My field' والی قطار بھریں۔"])
    guide.append(["2. فصل اور فصل کی حالت کے لیے فہرست سے انتخاب کریں۔ باقی خانوں میں عدد لکھیں؛ خالی خانہ نہ چھوڑیں۔"])
    guide.append(["3. فائل محفوظ کریں، FieldOps AI میں Excel فائل اپ لوڈ کریں، اپنی قطار منتخب کریں، پھر جانچ چلائیں۔"])
    guide.append(["4. چار تیار نمونے بھی ہیں: خشک گندم، بارش، گرمی/ہوا، زیادہ نمی۔ ان سے نتیجے کا فرق دیکھیں۔"])
    guide.append(["5. یہ نمونے صرف سافٹ ویئر ٹیسٹ ہیں۔ کھیت کے اصل فیصلے مقامی زرعی ماہر سے چیک کریں۔"])
    guide.append([""])
    guide.append(["Fill only the yellow 'My field' row, save, upload this workbook, select the row, and run the field check."])
    guide.column_dimensions["A"].width = 115
    for cell in guide[1]:
        cell.fill = header_fill
        cell.font = header_font
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    for row_number in range(2, guide.max_row + 1):
        guide.row_dimensions[row_number].height = 35

    photos = workbook.create_sheet("Sample Photos")
    photos.sheet_view.rightToLeft = True
    photos.column_dimensions["A"].width = 62
    photos.append(["مشقی فصل کی تصاویر / Synthetic upload-test pictures"])
    photos.append(["یہ کمپیوٹر سے بنی تصویریں ہیں، اصلی فصل نہیں۔ صرف اپ لوڈ ٹیسٹ کے لیے استعمال کریں۔"])
    photos["A1"].fill = header_fill
    photos["A1"].font = header_font
    photos["A1"].alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    photos["A2"].font = Font(color="9C3E2D", bold=True, size=12)
    photos["A2"].alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    photos.row_dimensions[1].height = 36
    photos.row_dimensions[2].height = 42
    image_dir = Path(__file__).resolve().parents[1] / "data" / "sample_images"
    sample_photos = [
        ("گندم / Wheat", "synthetic_wheat_upload_test.jpg"),
        ("کپاس / Cotton", "synthetic_cotton_upload_test.jpg"),
        ("چاول / Rice", "synthetic_rice_upload_test.jpg"),
    ]
    for index, (label, filename) in enumerate(sample_photos):
        label_row = 4 + index * 15
        photos.cell(label_row, 1, label)
        photos.cell(label_row, 1).font = Font(bold=True, size=14, color="28583E")
        image_path = image_dir / filename
        if image_path.exists():
            image = SpreadsheetImage(str(image_path))
            image.width = 480
            image.height = 320
            photos.add_image(image, f"A{label_row + 1}")
            photos.row_dimensions[label_row + 1].height = 180

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _find_columns(values: tuple[object, ...]) -> dict[str, int]:
    columns = {}
    for column, value in enumerate(values):
        key = NORMALIZED_ALIASES.get(_normalize_header(value))
        if key is not None:
            columns[key] = column
    return columns


def _number(value: object, key: str, errors: list[str]) -> float | None:
    if value is None or str(value).strip() == "":
        errors.append(FIELD_LABELS_UR[key])
        return None
    try:
        normalized = str(value).translate(URDU_DIGITS).replace(",", "").strip()
        number = float(normalized)
        if not math.isfinite(number):
            raise ValueError
        return number
    except (TypeError, ValueError):
        errors.append(FIELD_LABELS_UR[key])
        return None


def _map_choice(value: object, aliases: dict[str, str]) -> str | None:
    if value is None:
        return None
    return aliases.get(str(value).strip().casefold())


def _parse_scenario(sheet_name: str, row_number: int, row: tuple[object, ...], columns: dict[str, int]) -> SpreadsheetScenario | None:
    raw_values = {key: row[column] if column < len(row) else None for key, column in columns.items()}
    if all(value is None or str(value).strip() == "" for key, value in raw_values.items() if key != "scenario"):
        scenario_name = str(raw_values.get("scenario") or "").strip()
        if not scenario_name:
            return None

    name = str(raw_values.get("scenario") or f"{sheet_name}, row {row_number}").strip()
    crop_aliases = {
        "wheat": "Wheat", "گندم": "Wheat",
        "cotton": "Cotton", "کپاس": "Cotton",
        "rice": "Rice", "چاول": "Rice",
        "maize": "Maize", "corn": "Maize", "مکئی": "Maize",
    }
    stage_aliases = {
        "seedling": "Seedling", "ابتدائی پودا": "Seedling", "نرسری": "Seedling",
        "vegetative": "Vegetative", "بڑھوتری": "Vegetative",
        "flowering": "Flowering", "flowering / heading": "Flowering", "پھول یا بالی": "Flowering", "پھول": "Flowering", "بالی": "Flowering",
        "maturity": "Maturity", "near maturity": "Maturity", "فصل پکنے کے قریب": "Maturity",
    }
    errors: list[str] = []
    crop = _map_choice(raw_values.get("crop"), crop_aliases)
    stage = _map_choice(raw_values.get("stage"), stage_aliases)
    if crop is None:
        errors.append("فصل")
    if stage is None:
        errors.append("فصل کی حالت")
    numbers = {
        key: _number(raw_values.get(source_key), key, errors)
        for key, source_key in (
            ("area", "area"),
            ("moisture", "moisture"),
            ("recent_rain", "recent_rain"),
            ("forecast_rain", "forecast_rain"),
            ("temperature", "temperature"),
            ("wind", "wind"),
        )
    }
    area = numbers["area"]
    moisture = numbers["moisture"]
    recent_rain = numbers["recent_rain"]
    forecast_rain = numbers["forecast_rain"]
    temperature = numbers["temperature"]
    wind = numbers["wind"]
    if area is not None and area <= 0:
        errors.append("رقبہ صفر سے زیادہ ہونا چاہیے")
    if moisture is not None and not 0 <= moisture <= 100:
        errors.append("مٹی کی نمی 0 سے 100 کے درمیان ہونی چاہیے")
    if recent_rain is not None and recent_rain < 0 or forecast_rain is not None and forecast_rain < 0:
        errors.append("بارش کی مقدار منفی نہیں ہو سکتی")
    if temperature is not None and not -10 <= temperature <= 65:
        errors.append("درجہ حرارت -10 سے 65 کے درمیان ہونا چاہیے")
    if wind is not None and not 0 <= wind <= 200:
        errors.append("ہوا کی رفتار 0 سے 200 کے درمیان ہونی چاہیے")

    if errors:
        return SpreadsheetScenario(name, sheet_name, row_number, None, ", ".join(dict.fromkeys(errors)))

    field = FieldAssessment(
        crop=crop,
        crop_stage=stage,
        field_area_acres=area,
        soil_moisture_pct=moisture,
        recent_rain_mm=recent_rain,
        forecast_rain_mm=forecast_rain,
        temperature_c=temperature,
        wind_kph=wind,
        location=str(raw_values.get("location") or "").strip(),
        notes=str(raw_values.get("notes") or "").strip(),
    )
    return SpreadsheetScenario(name, sheet_name, row_number, field)


def read_workbook(content: bytes) -> list[SpreadsheetScenario]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    scenarios = []
    for sheet in workbook.worksheets:
        first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        columns = _find_columns(tuple(first_row))
        if not {"crop", "stage", "area", "moisture", "recent_rain", "forecast_rain", "temperature", "wind"}.issubset(columns):
            continue
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            scenario = _parse_scenario(sheet.title, row_number, tuple(row), columns)
            if scenario is not None:
                scenarios.append(scenario)
    workbook.close()
    if not scenarios:
        raise ValueError("Excel file needs a row of headers and at least one field row. Use the FieldOps template.")
    return scenarios